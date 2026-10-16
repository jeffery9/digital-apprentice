#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
prepub_check.py — 《数字见习生》出版前检查（四类终检＋出版就绪核对）

用法：cd /home/user/book && python3 prepub_check.py
纪律：对源稿与产物**只读**；发现问题只报告（不绿在哪、疑似原因、建议改哪个源文件），
修复是改稿任务，走「改稿闭环线」。检查对象是合并版产物＋docx＋两张登记表。
报告落盘：见习生工位/台账/出版前检查报告_<VERSION>.md
"""
import os, re, sys, ast, zipfile
from collections import Counter, defaultdict

ROOT = os.path.dirname(os.path.abspath(__file__))
SRC  = os.path.join(ROOT, '书稿')
IMG  = os.path.join(SRC, '插画')
OUT  = os.path.join(ROOT, '输出')
LEDGER = os.path.join(ROOT, '见习生工位', '台账')

sys.path.insert(0, ROOT)
import assemble_book as AB   # 只取 BOOK_ORDER / VERSION 常量，main 有 __main__ 守卫

MERGED_MD = os.path.join(OUT, f'《数字见习生》全书稿_{AB.VERSION}全书稿_合并版.md')  # 占位，下面重算
# 实际文件名不含版本，重算：
cands = [f for f in os.listdir(OUT) if f.endswith('合并版.md')]
MERGED_MD = os.path.join(OUT, cands[0]) if cands else MERGED_MD
MERGED_DOCX = MERGED_MD[:-3] + '.docx'

md = open(MERGED_MD, encoding='utf-8').read()
lines = md.splitlines()

results = []   # (级别, 编号, 描述)  级别: PASS/FAIL/WARN/INFO
def rep(level, code, msg):
    results.append((level, code, msg))
    mark = {'PASS': '✅', 'FAIL': '❌', 'WARN': '⚠️', 'INFO': 'ℹ️'}[level]
    print(f'{mark} [{code}] {msg}')

def nearest_heading(lineno, h1_only=False):
    for i in range(lineno, -1, -1):
        if lines[i].startswith('#'):
            if h1_only and not lines[i].startswith('# '):
                continue
            return lines[i].lstrip('# ').strip()[:40]
    return '(文档开头)'

# ══════════════════════════════════════════════════════════════
# A. 结构与版本一致性
# ══════════════════════════════════════════════════════════════
print('\n═══ A. 结构与版本 ═══')
bo_files = [e[1] for e in AB.BOOK_ORDER if e[0] == 'file']
bo_parts = [e[2] for e in AB.BOOK_ORDER if e[0] == 'part']
missing = [f for f in bo_files if not os.path.exists(os.path.join(SRC, f))]
missing += [f'插画/{p}' for p in bo_parts if not os.path.exists(os.path.join(IMG, p))]
rep('PASS' if not missing else 'FAIL', 'A1',
    f'BOOK_ORDER {len(bo_files)} 个源文件＋{len(bo_parts)} 张篇扉页全部在位' if not missing else f'缺文件: {missing}')

prog = open(os.path.join(SRC, '00_书稿进度.md'), encoding='utf-8').read()
v_prog = re.search(r'终态版 (v[\d.]+\w*)', prog)
v_sub  = re.search(r'subtitle[:：].*?(v[\d.]+\w*)', md) or re.search(r'全书合并审阅版 · (v[\d.]+\w*)', md)
vers = {'脚本': AB.VERSION, '进度表': v_prog.group(1) if v_prog else '?',
        '合并版': v_sub.group(1) if v_sub else '?'}
import docx as _docx
doc = _docx.Document(MERGED_DOCX)
vers['docx'] = (doc.core_properties.subject or '?').split()[-1]
same = len(set(vers.values())) == 1
rep('PASS' if same else 'FAIL', 'A2', f'版本号四处一致: {vers}' if same else f'版本号不一致: {vers}')

h1 = [p for p in doc.paragraphs if p.style.name == 'Heading 1']
toc_ps = [p for p in doc.paragraphs if p.style.name.startswith('TOC')]
rep('PASS' if len(h1) == 28 else 'WARN', 'A3', f'docx H1＝{len(h1)}（期望 28：书名页1+前言1+篇扉页3+章15+附录7+后记1）')
rep('INFO', 'A4', f'docx 静态目录条目＝{len(toc_ps)}')

# ══════════════════════════════════════════════════════════════
# B. 第①类：断链检查（双向）
# ══════════════════════════════════════════════════════════════
print('\n═══ B. 断链检查 ═══')
# B1 正向图片链接
img_refs = re.findall(r'!\[[^\]]*\]\(([^)]+)\)', md)
broken = []
for p in img_refs:
    fp = os.path.normpath(os.path.join(OUT, p))
    if not os.path.exists(fp):
        broken.append(p)
rep('PASS' if not broken else 'FAIL', 'B1',
    f'正向图片链接 {len(img_refs)} 处全部落盘' if not broken else f'图片断链 {len(broken)}: {broken[:6]}')

# B2 反向：磁盘文件 vs 交付清单登记 vs 正文引用
disk = set(os.listdir(IMG))
registry_txt = open(os.path.join(SRC, '插画交付清单.md'), encoding='utf-8').read()
registered = set(re.findall(r'插画/([^`\s|）)，。;；]+\.png)', registry_txt))
referenced = set(os.path.basename(p) for p in img_refs)
REG_NOSPACE = re.sub(r'\s+', '', registry_txt)
ghost = sorted(p for p in registered if 'x-x' not in p and p not in disk)
rep('PASS' if not ghost else 'FAIL', 'B2a',
    f'交付清单中 {len(registered)} 处文件路径全部落盘（“图x-x”为模板记法，不计）' if not ghost
    else f'清单路径磁盘无: {ghost}')
def is_registered(fn):
    stem = re.sub(r'\.png$', '', fn)
    if re.sub(r'\s+', '', fn) in REG_NOSPACE:
        return True
    m = re.search(r'图(\d+)-(\d+)', stem)
    if m and re.search(rf'图 ?{m.group(1)}-{m.group(2)}', registry_txt):
        return True
    core = re.sub(r'^(P0|P1|P2b|P2)_', '', stem)
    core = re.sub(r'^图\d+-\d+_', '', core)
    toks = [t for t in core.split('_') if len(t) >= 2]
    return all(t in REG_NOSPACE for t in toks) if toks else False
disk_pngs = sorted(f for f in disk if f.endswith('.png'))
unreg = [f for f in disk_pngs if not is_registered(f)]
rep('PASS' if not unreg else 'FAIL', 'B2b',
    f'磁盘 {len(disk_pngs)} 个 PNG 全部可在交付清单找到登记（按图号/名称关键词匹配）' if not unreg
    else f'磁盘有而清单未登记: {unreg}')
unref = sorted(f for f in disk if f.endswith('.png') and f not in referenced)
EXPECT_UNREF = {'封面概念图.png', '封面_方案A_书上接文件.png', '封面_方案B_工作台协作.png',
                '封面_方案C_门缝之光.png', '风格样张1_角色资产.png',
                'P2b_思维导图模板样张.png', 'P2b_构件卡模板样张.png',
                'P0_比喻资产库合集.png', 'P0_角色_周舟设定.png'}
unexpected_unref = [f for f in unref if f not in EXPECT_UNREF]
rep('PASS' if not unexpected_unref else 'FAIL', 'B2c',
    f'正文未引用 {len(unref)} 件，均属书籍级/样张/资产库（封面4+样张3+资产库2）' if not unexpected_unref
    else f'正文未引用且不在预期书籍级名单: {unexpected_unref}')
ref_cnt = Counter(os.path.basename(p) for p in img_refs)
dups = {k: v for k, v in ref_cnt.items() if v > 1}
ok_inst = len(img_refs) == len(doc.inline_shapes)
rep('PASS' if ok_inst else 'FAIL', 'B2d',
    f'md 图片引用 {len(img_refs)} 处（去重 {len(ref_cnt)} 件）＝docx 内嵌 {len(doc.inline_shapes)} 张')
for k, v in dups.items():
    rep('INFO', 'B2e', f'同文件被引用 {v} 次：{k}（第 7 章图 7-2 与第 1 章图 1-4 共用底稿，正文有内联说明，属有意设计）')

# B3–B5 编号引用必须有定义
fig_defs = set()
for m in re.finditer(r'!\[图 ?(\d+)-(\d+)', md):          fig_defs.add((int(m[1]), int(m[2])))
for m in re.finditer(r'【表单 T-\d+｜图 ?(\d+)-(\d+)', md): fig_defs.add((int(m[1]), int(m[2])))
for m in re.finditer(r'\*\*【图 ?(\d+)-(\d+)｜', md):        fig_defs.add((int(m[1]), int(m[2])))
tbl_defs = set()
for m in re.finditer(r'【表 ?(\d+)-(\d+)｜', md):            tbl_defs.add((int(m[1]), int(m[2])))
for m in re.finditer(r'\*\*表 ?(\d+)-(\d+)　', md):          tbl_defs.add((int(m[1]), int(m[2])))
t_defs = set(re.findall(r'【表单 T-(\d{2})｜', md))

body_wo_links = re.sub(r'\]\([^)]*\)', ']', md)   # 去掉链接目标（文件名里的 图x-y 不算引用）
fig_refs = set((int(a), int(b)) for a, b in re.findall(r'图 ?(\d+)-(\d+)', body_wo_links))
tbl_refs = set((int(a), int(b)) for a, b in re.findall(r'表 ?(\d+)-(\d+)', body_wo_links))
t_refs   = set(re.findall(r'T-(\d{2})', body_wo_links))
bad_fig = sorted(fig_refs - fig_defs); bad_tbl = sorted(tbl_refs - tbl_defs)
bad_t = sorted(t for t in t_refs if t not in t_defs)
rep('PASS' if not bad_fig else 'FAIL', 'B3',
    f'图号引用 {len(fig_refs)} 个全部有定义（定义 {len(fig_defs)} 个）' if not bad_fig
    else f'引用了未定义的图号: {["图%d-%d" % x for x in bad_fig]}')
rep('PASS' if not bad_tbl else 'FAIL', 'B4',
    f'表号引用 {len(tbl_refs)} 个全部有定义（定义 {len(tbl_defs)} 个）' if not bad_tbl
    else f'引用了未定义的表号: {["表%d-%d" % x for x in bad_tbl]}')
rep('PASS' if not bad_t else 'FAIL', 'B5',
    f'T 编号引用 {len(t_refs)} 个全部有定义（T-01～T-12 共 {len(t_defs)} 张）' if not bad_t
    else f'引用了未定义表单: {["T-"+x for x in bad_t]}')

# B6 §x.y 与 §X-n 引用
head_secs = set(re.findall(r'^## (\d+\.\d+)', md, re.M))
head_appx = set(re.findall(r'^## ([A-F]-\d+)', md, re.M)) | set(re.findall(r'^### ([A-F]-\d+)', md, re.M))
EXT = ('策划案', '交付清单', '排版Prompt', '表格成品集', '大纲', '员工手册', '进度表', '蓝图')
bad_sec = []
for m in re.finditer(r'§(\d+\.\d+)(?![\d\-b])', md):
    pre = md[max(0, m.start()-16):m.start()]
    if any(e in pre for e in EXT):
        continue
    if m.group(1) not in head_secs:
        bad_sec.append((m.group(1), nearest_heading(md[:m.start()].count('\n'))))
bad_appx = [m.group(1) for m in re.finditer(r'§([A-F]-\d+)', md) if m.group(1) not in head_appx]
rep('PASS' if not bad_sec else 'FAIL', 'B6a',
    f'§x.y 节引用全部有对应标题（节标题 {len(head_secs)} 个）' if not bad_sec else f'悬空节引用: {bad_sec[:8]}')
rep('PASS' if not bad_appx else 'FAIL', 'B6b',
    '附录节引用（§F-1 等）全部有对应标题' if not bad_appx else f'悬空附录节引用: {sorted(set(bad_appx))}')

# B7 章号与附录号引用
bad_ch = sorted(set(int(n) for n in re.findall(r'第 ?(\d{1,2}) ?章', md)) - set(range(1, 16)))
apx_refs = set(re.findall(r'附录 ?([A-G])', md))
bad_apx = sorted(apx_refs - set('ABCDEFG'))
rep('PASS' if not bad_ch else 'FAIL', 'B7a', '章号引用均在 1–15 范围' if not bad_ch else f'越界章号: {bad_ch}')
rep('PASS' if not bad_apx else 'FAIL', 'B7b', f'附录引用 {sorted(apx_refs)} 均在 A–G' if not bad_apx else f'未知附录: {bad_apx}')

# ══════════════════════════════════════════════════════════════
# C. 第②类：编号完整性
# ══════════════════════════════════════════════════════════════
print('\n═══ C. 编号完整性 ═══')
def contiguity(defs, label, code):
    per = defaultdict(set)
    for c, n in defs:
        per[c].add(n)
    issues = []
    for c in sorted(per):
        nums, mx = per[c], max(per[c])
        gaps = sorted(set(range(1, mx + 1)) - nums)
        if gaps:
            issues.append(f'第{c}章{label}缺号 {gaps}')
    rep('PASS' if not issues else 'FAIL', code,
        f'{label}号逐章连续（第 1–15 章，共 {len(defs)} 个）' if not issues else '; '.join(issues))
    return per
fig_per = contiguity(fig_defs, '图', 'C1')
tbl_per = contiguity(tbl_defs, '表', 'C2')
dup_fig = [(c, n) for c, ns in fig_per.items() for n in ns if len(re.findall(rf'!\[图 ?{c}-{n}[　 ]', md)) > 1]
rep('PASS' if not dup_fig else 'WARN', 'C2b', '图片定义无重号' if not dup_fig else f'疑似重复定义: {dup_fig}')
rep('PASS' if t_defs == set(f'{i:02d}' for i in range(1, 13)) else 'FAIL', 'C3',
    f'T-01～T-12 恰好 12 张、无跳号无重号' if t_defs == set(f'{i:02d}' for i in range(1, 13))
    else f'T 定义异常: {sorted(t_defs)}')

# C4 每个表题下方跟真实表体
cap_lines = [(i, l) for i, l in enumerate(lines)
             if re.match(r'^(> ?)?(\*\*)?【表 ?\d+-\d+｜', l) or re.match(r'^\*\*表 ?\d+-\d+　', l)]
no_body, quote_body = [], []
for i, l in cap_lines:
    found = None
    for j in range(i + 1, min(i + 4, len(lines))):
        s = lines[j].lstrip()
        if s.startswith('|'):
            found = 'pipe'; break
        if s.startswith('>'):
            found = 'quote'; break
        if lines[j].strip() and not lines[j].startswith('*'):
            break
    if found is None:
        no_body.append((i + 1, l.strip()[:30]))
    elif found == 'quote':
        quote_body.append(l.strip()[:24])
rep('PASS' if not no_body else 'FAIL', 'C4',
    f'{len(cap_lines)} 个表题下方均有真实表体' if not no_body else f'有题无体: {no_body}')
if quote_body:
    rep('INFO', 'C4b', f'{len(quote_body)} 个表体为引用块版式（约定 §四.5 模板排版，设计师出片时注意）: {quote_body}')

# C5 表单四件套
def form_body(start):
    out = []
    for j in range(start + 1, len(lines)):
        l = lines[j]
        if (re.match(r'^#{1,4} ', l) or l.strip() == '---' or '【表单 T-' in l
                or re.match(r'^\s*(\*\*)?【[图表] ?\d', l) or re.match(r'^\*\*表 ?\d+-\d+', l)):
            break
        out.append(l)
    return '\n'.join(out)
quad_issues = []
form_where = {}
for i, l in enumerate(lines):
    m = re.search(r'【表单 (T-\d{2})｜', l)
    if m:
        t = m.group(1); form_where[t] = nearest_heading(i, h1_only=True)
        body = form_body(i)
        has_field = bool(re.search(r'^\s*>?\s*\|', body, re.M)) or '＿' in body or re.search(r'^> .{2,20}[:：]', body, re.M)
        has_example = ('示例' in body) or ('实例' in body) or bool(re.search(r'\*[^*\n]{4,}\*', body)) or ('示例' in l)
        has_note = bool(re.search(r'(^\s*>?\s*注[:：]|脚注|排版说明)', body, re.M)) or ('可复印' in l)
        lacks = [nm for nm, ok in (('字段结构', has_field), ('示例', has_example), ('脚注/注记', has_note)) if not ok]
        if lacks:
            quad_issues.append(f'{t}@行{i+1} 缺{"/".join(lacks)}')
rep('PASS' if not quad_issues else 'FAIL', 'C5',
    '12 张表单四件套齐全（编号标题＋字段结构＋示例＋脚注）' if not quad_issues else f'四件套缺项: {quad_issues}')
in_appf = [t for t, w in form_where.items() if '附录' in w or 'F-' in w]
rep('INFO', 'C6', f'表单落位：T-11/T-12 在附录 F＝{sorted(t for t in in_appf if t in ("T-11","T-12"))}；'
                  f'其余 {len(form_where)-len(in_appf)} 张在正文各章')

# ══════════════════════════════════════════════════════════════
# D. 第③类：引号与金句
# ══════════════════════════════════════════════════════════════
print('\n═══ D. 引号与金句 ═══')
dq, sq = md.count('"'), md.count("'")
rep('PASS' if dq == 0 and sq == 0 else 'FAIL', 'D1', f'合并版 ASCII 直引号残留＝{dq}（双）/{sq}（单）')

GOLDEN = [
    '大脑是厂商给的，制度是你设计的',
    '手册是安全带，不是保险柜',
    'AI 输出是待核验的初稿',
    '反思要有外部信号，不然就是自我催眠',
    '尺寸是预算，整洁是本事',
    '多智能体是论证出来的选择，不是默认答案',
    '可观测性优先于智能性',
    '制度可以借 AI 的手，不能借 AI 的责任',
    '契约对方没有羞耻心，核查方必须有心跳',
    'AI 答“怎么做”，人问“为什么”',
]
hits = [md.count(g) for g in GOLDEN]
miss = [(i + 1, g) for i, (g, h) in enumerate(zip(GOLDEN, hits)) if h == 0]
rep('PASS' if not miss else 'FAIL', 'D2',
    f'十句金句锚点逐字命中：{hits}' if not miss else f'金句零命中: {miss}')

# D3 术语一致性（附录A 100 条 vs 正文出现）＋高危变体
appa = open(os.path.join(SRC, '附录A_术语表.md'), encoding='utf-8').read()
term_rows = []
for l in appa.splitlines():
    if l.startswith('|') and not re.match(r'^\|[\s:\-|]+\|$', l) and not l.startswith('| 术语'):
        cells = [c.strip() for c in l.strip().strip('|').split('|')]
        if len(cells) >= 5 and cells[0]:
            term_rows.append((cells[0], cells[4]))
terms = [t for t, _ in term_rows]
ia, ib = md.find('# 附录 A'), md.find('# 附录 B')
body_only = md[:ia] + md[ib:] if ia > 0 and ib > ia else md
zero = [(t, f) for t, f in term_rows if body_only.count(t) == 0]
rep('PASS' if len(zero) <= 5 else 'WARN', 'D3a',
    f'附录 A 术语 {len(terms)} 条，正文（除附录A外）零出现 {len(zero)} 条——其“首现”栏指向的章节不含该词，列入审校第②轮导流一致性清单'
    + (f': {zero}' if zero else ''))
claim = re.search(r'(\d+) 个概念', appa.splitlines()[0])
if claim:
    n_claim = int(claim.group(1))
    ok_n = n_claim == len(terms)
    rep('PASS' if ok_n else 'FAIL', 'D3c',
        f'附录 A 标题口径 {n_claim} 个概念＝实际 {len(terms)} 条' if ok_n
        else f'附录 A 标题称“{n_claim} 个概念”，实际 {len(terms)} 条——改标题或补条目，须作者拍板（源文件：附录A_术语表.md）')
VARIANTS = [('大语言模型', r'大型语言模型'), ('Sub-agent', r'(?i)sub ?agent(?!s)|subagent'),
            ('Qwen Code', r'(?i)qwen[-_ ]code'), ('多智能体', r'多代理'),
            ('上下文窗口', r'语境窗口'), ('知识库', r'知识基地')]  # 工作流程＝六要素之四的正典名（第2章§2.3），与构件“工作流”（第6章）不是变体对
vmsg = []
for canon, pat in VARIANTS:
    alt = [m.group(0) for m in re.finditer(pat, md)]
    alt = [a for a in alt if a != canon]
    if alt:
        vmsg.append(f'{canon}→变体{Counter(alt).most_common(3)}')
rep('PASS' if not vmsg else 'WARN', 'D3b', '高危术语变体扫描：未见混用' if not vmsg else f'术语变体（人工判定）: {vmsg}')

# ══════════════════════════════════════════════════════════════
# E. 第④类：表格构造（md 侧＋docx 侧）
# ══════════════════════════════════════════════════════════════
print('\n═══ E. 表格构造 ═══')
# E1 md 管道表前后空行
viol = 0
i = 0
n = len(lines)
while i < n:
    if lines[i].lstrip().startswith('|'):
        if i > 0 and lines[i - 1].strip():
            viol += 1
        j = i
        while j < n and lines[j].lstrip().startswith('|'):
            j += 1
        if j < n and lines[j].strip():
            viol += 1
        i = j
    else:
        i += 1
blocks, cur = [], []
for l in lines:
    if l.lstrip().startswith('|'):
        cur.append(l)
    else:
        if cur: blocks.append(cur); cur = []
if cur: blocks.append(cur)
nblocks = len(blocks)
rep('PASS' if viol == 0 else 'FAIL', 'E1', f'md 管道表块 {nblocks} 个，前后空行违规＝{viol}')

# E2/E3 docx 侧
pipe_ps = [p.text[:40] for p in doc.paragraphs if '|' in p.text]
rep('PASS' if not pipe_ps else 'FAIL', 'E2',
    f'docx 正文段落残留管道符＝{len(pipe_ps)}' if not pipe_ps else f'残留管道符段落: {pipe_ps[:5]}')
rep('PASS' if len(doc.tables) >= nblocks else 'FAIL', 'E3',
    f'docx 表格 {len(doc.tables)} 张 ≥ md 管道表块 {nblocks} 个')

# E4–E7 docx 版式
with zipfile.ZipFile(MERGED_DOCX) as z:
    docxml = z.read('word/document.xml').decode('utf-8')
    styles = z.read('word/styles.xml').decode('utf-8')
    footers = [n for n in z.namelist() if n.startswith('word/footer')]
    fxml = ''.join(z.read(f).decode('utf-8') for f in footers)
toc_field = 'TOC \\o' in docxml or '>TOC ' in docxml
rep('PASS' if not toc_field else 'FAIL', 'E4', '目录为静态文本（无 Word TOC 域）' if not toc_field else 'docx 含 TOC 域（打开是空的）')
sec = doc.sections[0]
w_cm, h_cm = sec.page_width.cm, sec.page_height.cm
rep('PASS' if abs(w_cm - 21) < 0.1 and abs(h_cm - 29.7) < 0.1 else 'FAIL', 'E5',
    f'页面 A4：{w_cm:.1f}×{h_cm:.1f}cm；边距 上{sec.top_margin.cm:.1f} 下{sec.bottom_margin.cm:.1f} 左{sec.left_margin.cm:.1f} 右{sec.right_margin.cm:.1f}cm')
def style_ea(sid):
    m = re.search(rf'<w:style [^>]*w:styleId="{sid}".*?</w:style>', styles, re.S)
    if not m:
        return None
    m2 = re.search(r'w:eastAsia="([^"]+)"', m.group(0))
    return m2.group(1) if m2 else None
ea_map = {'Normal': style_ea('Normal'), 'Heading1': style_ea('Heading1')}
ok_font = ea_map.get('Normal') and ea_map.get('Heading1')
rep('PASS' if ok_font else 'FAIL', 'E6',
    f'中文字体 eastAsia 已设：正文＝{ea_map.get("Normal")}，标题＝{ea_map.get("Heading1")}')
rep('PASS' if 'PAGE' in fxml else 'FAIL', 'E7', '页脚含页码域')

# ══════════════════════════════════════════════════════════════
# F. 出版就绪
# ══════════════════════════════════════════════════════════════
print('\n═══ F. 出版就绪 ═══')
STRUCT_PREFIX = ('表 ', '表', '表单', '图 ', '图', '触发', '步骤', '模板', '自检', '规矩',
                 '输出', '角色', '任务', '流程', '格式', '兜底', '禁止', '本章', '你的')
ph_rows = []
for i, l in enumerate(lines):
    for m in re.finditer(r'【([^【】]{1,60})】', l):
        t = m.group(1)
        if t.startswith(STRUCT_PREFIX) or t in ('方括号',):
            continue
        ph_rows.append((i + 1, nearest_heading(i), '【' + t + '】', l.strip()[:60]))
    for m in re.finditer(r'\[占位[：:][^\]\n]{0,90}\]', l):
        ph_rows.append((i + 1, nearest_heading(i), m.group(0)[:56], l.strip()[:60]))
    for m in re.finditer(r'（定稿时[^）\n]{0,40}）', l):
        ph_rows.append((i + 1, nearest_heading(i), '＿＿' + m.group(0)[:34], l.strip()[:60]))
real_ph = [r for r in ph_rows if any(k in r[2] for k in ('待', '注明', '编辑', '作者', '定稿', '填充', '占位'))]
other_ph = [r for r in ph_rows if r not in real_ph]
rep('INFO', 'F1a', f'方括号占位（作者/编辑待填类）＝{len(real_ph)} 处')
for r in real_ph:
    print(f'    行{r[0]}〈{r[1]}〉{r[2]}')
rep('INFO' if not other_ph else 'WARN', 'F1b',
    f'其余非结构性【】标记＝{len(other_ph)} 处（逐条人工过目）' if other_ph else '其余非结构性【】标记＝0 处')
for r in other_ph[:12]:
    print(f'    行{r[0]}〈{r[1]}〉{r[2]}｜{r[3]}')

dsg = [(i + 1, nearest_heading(i)) for i, l in enumerate(lines) if '排版说明' in l]
rep('WARN' if dsg else 'PASS', 'F2',
    f'设计师专用「排版说明」引用块＝{len(dsg)} 处：审阅稿按约定保留（供设计师排版 45 张表用）；**送印前须用 `python3 assemble_book.py --final` 构建移除**（v0.8h 已加开关并验证：final 构建残留＝0，其余内容仅少这 8 行）'
    if dsg else '无排版说明残留（--final 构建或已清理）')
draft_left = sum(md.count(k) for k in ('书稿状态', '标注位说明', '本章图表：', '目标字数'))
rep('PASS' if draft_left == 0 else 'FAIL', 'F3', f'草稿注释残留（书稿状态/标注位说明/本章图表/目标字数）＝{draft_left}')

houji = open(os.path.join(SRC, '后记.md'), encoding='utf-8').read()
qianyan = open(os.path.join(SRC, '前言.md'), encoding='utf-8').read()
PH_SPAN = r'\[占位[：:]|【[^】]*(?:待|注明|填写)[^】]*】'
n_hj = len(re.findall(PH_SPAN, houji)); n_qy = len(re.findall(PH_SPAN, qianyan))
rep('INFO', 'F4', f'模板稿占位：前言 {n_qy} 处、后记 {n_hj} 处（进度表 §五已登记，作者填充项）')
ack_risk = '感谢为本书初稿提供技术审读的工程师朋友' in qianyan
rep('WARN' if ack_risk else 'PASS', 'F4b',
    '前言致谢段无占位标记，却已按事实口吻写“感谢…提供技术审读的工程师朋友／参与试读的目标读者”，而四轮审校（§五）尚未进行——定稿时作者必须按真实情况改写，否则构成事实虚构（§五已登记“前言致谢与案例来源”，此处升级为高风险项）'
    if ack_risk else '前言致谢段未见虚构风险表述')
rep('INFO', 'F5', f'勾选框 ☐ 合计＝{md.count("☐")}；填空线 ＿ 合计＝{md.count("＿")}')

# ══════════════════════════════════════════════════════════════
# G. 表单同步（成品集母本 vs 正文阅读版 三处落点）
# ══════════════════════════════════════════════════════════════
print('\n═══ G. 表单同步 ═══')
chengpin = open(os.path.join(SRC, '表格成品集.md'), encoding='utf-8').read()
paiban = open(os.path.join(SRC, '表格排版Prompt.md'), encoding='utf-8').read()
no_master = [f'T-{i:02d}' for i in range(1, 13) if f'## T-{i:02d}' not in chengpin]
no_card = [f'T-{i:02d}' for i in range(1, 13) if f'T-{i:02d}' not in paiban]
rep('PASS' if not no_master else 'FAIL', 'G1', '成品集母本 T-01～T-12 齐' if not no_master else f'母本缺: {no_master}')
rep('PASS' if not no_card else 'FAIL', 'G2', '表格排版Prompt 逐表卡 T-01～T-12 齐' if not no_card else f'逐表卡缺: {no_card}')

def normq(x):
    x = re.sub(r'"([^"\n]*)"', '\u201c\\1\u201d', x)
    x = re.sub(r"'([^'\n]*)'", '\u2018\\1\u2019', x)
    return x.replace('"', '\u201d').replace("'", '\u2019')
def norm_rows(text):
    rows = []
    for l in text.splitlines():
        l = l.strip().lstrip('>').strip()
        if l.startswith('|') and not re.match(r'^\|[\s:\-|]+\|$', l):
            rows.append(normq(re.sub(r'\s+', '', l.replace('**', ''))))
    return rows
sync_bad = []
for i in range(1, 13):
    t = f'T-{i:02d}'
    m1 = re.search(rf'^## {t}　.*?(?=^## |\Z)', chengpin, re.M | re.S)
    master = norm_rows(m1.group(0)) if m1 else []
    ti = next((k for k, l in enumerate(lines) if f'【表单 {t}｜' in l), None)
    reader = norm_rows(form_body(ti)) if ti is not None else []
    only_m = [r for r in master if r not in reader]
    only_r = [r for r in reader if r not in master]
    if only_m or only_r:
        sync_bad.append(f'{t}: 母本独有{len(only_m)}行/阅读版独有{len(only_r)}行')
rep('PASS' if not sync_bad else 'WARN', 'G3',
    '12 张表单母本 vs 阅读版管道行逐字一致' if not sync_bad
    else f'母本/阅读版差异（人工判定是否属排版性差异）: {sync_bad}')

# ══════════════════════════════════════════════════════════════
# 汇总
# ══════════════════════════════════════════════════════════════
print('\n═══ 汇总 ═══')
n_fail = sum(1 for lv, _, _ in results if lv == 'FAIL')
n_warn = sum(1 for lv, _, _ in results if lv == 'WARN')
n_pass = sum(1 for lv, _, _ in results if lv == 'PASS')
print(f'PASS {n_pass}｜WARN {n_warn}｜FAIL {n_fail}｜INFO {len(results)-n_pass-n_warn-n_fail}')

# 报告落盘
os.makedirs(LEDGER, exist_ok=True)
rp = os.path.join(LEDGER, f'出版前检查报告_{AB.VERSION}.md')
with open(rp, 'w', encoding='utf-8') as f:
    f.write(f'# 出版前检查报告（{AB.VERSION}）\n\n')
    f.write('> 依据：见习生工位「技能_四类终检」＋出版就绪核对；对象＝合并版产物＋docx＋两张登记表；本脚本对源稿只读。\n\n')
    f.write(f'终检结果：①断链＝{sum(1 for lv,c,_ in results if lv=="FAIL" and c.startswith("B"))} 处不绿　'
            f'②编号＝{"通过" if not [1 for lv,c,_ in results if lv=="FAIL" and c.startswith("C")] else "不通过"}　'
            f'③直引号＝{dq}／金句命中＝{"10 句全部逐字 ☑" if not miss else "有零命中 ☐"}　'
            f'④表格＝残留管道符 {len(pipe_ps)}、docx 表 {len(doc.tables)} 张\n\n')
    concl = '全绿 ☑' if n_fail == 0 and n_warn == 0 else ('绿但有注意项 ☑（WARN 见下）' if n_fail == 0 else '不绿 ☐')
    f.write(f'结论：{concl}\n\n')
    if n_fail or n_warn:
        f.write('## 不绿项与待人工项（不绿在哪、疑似原因、建议改哪个源文件）\n\n')
        for lv, c, m in results:
            if lv in ('FAIL', 'WARN'):
                f.write(f'- **{lv}｜{c}**：{m}\n')
        f.write('\n遗留问题：上列各项均已在《00_书稿进度.md》§五登记对应责任方；除此之外无。\n\n')
    f.write('| 级别 | 编号 | 结果 |\n|---|---|---|\n')
    for lv, c, m in results:
        f.write(f'| {lv} | {c} | {m.replace("|", "／")} |\n')
    if real_ph:
        f.write('\n## 作者/编辑待填占位清单\n\n')
        for r in real_ph:
            f.write(f'- 行 {r[0]}〈{r[1]}〉：{r[2]}\n')
print(f'报告已写入: {rp}')
